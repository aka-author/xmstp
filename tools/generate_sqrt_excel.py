"""Generate an Excel workbook with square roots for natural numbers 0..20."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parent.parent / "square_roots_0_to_20.xlsx"


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Square roots"

    ws["A1"] = "n"
    ws["B1"] = "√n"
    for cell in (ws["A1"], ws["B1"]):
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for n in range(0, 21):
        row = n + 2
        ws.cell(row=row, column=1, value=n)
        ws.cell(row=row, column=2, value=n**0.5)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
